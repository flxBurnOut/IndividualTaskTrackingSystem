"""Registered document executor. It cannot execute user-provided code."""
from __future__ import annotations
import json
import os
from pathlib import Path
import subprocess
import sys
import time
from xml.sax.saxutils import escape

from .resources import ResourceError, _plain_path, _relative


def produce(manager, job_id, value, cancel):
    meta = manager.create_workspace(job_id)
    workspace = Path(meta['path'])
    relative = _relative(value['relative_path'])
    target = _plain_path(workspace / relative)
    if target.exists():
        raise ResourceError('ALREADY_EXISTS', '成果文件已存在，请创建新的版本。')
    target.parent.mkdir(parents=True, exist_ok=True)
    manager._update_workspace(job_id, {'relative_path': relative, 'kind': value['kind'], 'state': 'registered', 'pinned': True})
    work = {**value, 'output': relative}
    if value['kind'] == 'pdf_notebook':
        work['source_path'] = str(manager._blob(value['source_sha256']))
    job_file = workspace / '.executor-input.json'
    with job_file.open('x', encoding='utf-8') as stream:
        json.dump(work, stream, ensure_ascii=False)
    args = [sys.executable]
    if getattr(sys, 'frozen', False):
        sibling = Path(sys.executable).with_name('PersonalManagementService.exe')
        args[0] = str(sibling if sibling.exists() else sys.executable)
        args += ['--document-worker', str(workspace)]
    else:
        args += ['-m', 'management.document_worker', str(workspace)]
    with manager._reservation(64 * 1024 * 1024):
        process = subprocess.Popen(args, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        deadline = time.monotonic() + 120
        import psutil
        try:
            while process.poll() is None:
                if cancel.is_set():
                    raise ResourceError('CANCELLED', '成果作业已取消，保留已登记的候选文件。')
                if time.monotonic() > deadline:
                    raise ResourceError('TIMEOUT', '文档处理超时；候选保留，可缩小范围后重试。')
                try:
                    parent = psutil.Process(process.pid)
                    members = [parent, *parent.children(recursive=True)]
                    private_bytes = sum(getattr(p.memory_info(), 'private', p.memory_info().rss) for p in members if p.is_running())
                    if private_bytes > 1024 ** 3:
                        raise ResourceError('MEMORY_LIMIT', '文档处理达到 1 GiB 预算，请分批处理。')
                except psutil.NoSuchProcess:
                    pass
                cancel.wait(.05)
            report_file = workspace / '.executor-result.json'
            if not report_file.exists():
                raise ResourceError('EXECUTOR_FAILED', '文档执行器未返回完整结果，文件仍保留。')
            report = json.loads(report_file.read_text('utf-8'))
            if not report.get('ok'):
                raise ResourceError(report.get('code', 'EXECUTOR_FAILED'), report.get('message', '文档生成失败。'))
            manager._update_workspace(job_id, {'relative_path': relative, 'kind': value['kind'], 'state': 'staged', 'validation': report['validation'], 'pinned': True})
            actual_kind = 'notebook' if value['kind'] == 'pdf_notebook' else 'file'
            frozen = manager.freeze_artifact(job_id, relative, kind=actual_kind, source_versions=value.get('source_versions', []))
            frozen['validation'].update(report['validation'])
            frozen['kind'] = value['kind']
            return frozen
        finally:
            if process.poll() is None:
                try:
                    parent = psutil.Process(process.pid)
                    for child in parent.children(recursive=True):
                        child.kill()
                    parent.kill()
                except psutil.NoSuchProcess:
                    pass
                process.wait(timeout=5)


def worker(workspace):
    workspace = _plain_path(workspace)
    request = json.loads((workspace / '.executor-input.json').read_text('utf-8'))
    result = {}
    try:
        output = workspace / _relative(request['output'])
        kind = request['kind']
        if kind == 'pdf_notebook':
            from pypdf import PdfReader
            source = _plain_path(request['source_path'], file=True)
            if source.stat().st_size > 256 * 1024 * 1024:
                raise ResourceError('RESOURCE_LIMIT', 'PDF 超过解析预算，请先分割。')
            reader = PdfReader(source)
            if reader.is_encrypted and not reader.decrypt(''):
                raise ResourceError('PASSWORD_REQUIRED', 'PDF 已加密，需要先提供可读取的版本。')
            pages = request['pages']
            if not pages or len(pages) > 20 or any(type(p) is not int or not 1 <= p <= len(reader.pages) for p in pages):
                raise ResourceError('INVALID_PAGES', '请选择存在的页码，每次最多 20 页。')
            cells = [{'cell_type': 'markdown', 'metadata': {}, 'source': '# ' + request.get('title', '练习') + '\n仅提取原文并留出作答区，未生成答案。'}]
            for page in pages:
                text = reader.pages[page - 1].extract_text() or ''
                if not text.strip():
                    raise ResourceError('OCR_REQUIRED', '第 %s 页没有可提取文本；需要 OCR 或人工转录，未生成空题。' % page)
                if len(text) > 60000:
                    raise ResourceError('RESOURCE_LIMIT', '单页文本过多，请缩小范围。')
                cells += [{'cell_type': 'markdown', 'metadata': {'source_page': page, 'source_sha256': request['source_sha256']}, 'source': '## 原文第 %s 页\n\n%s' % (page, text)}, {'cell_type': 'code', 'metadata': {}, 'source': '# 在此作答\n', 'execution_count': None, 'outputs': []}]
            notebook = {'nbformat': 4, 'nbformat_minor': 5, 'metadata': {'source_sha256': request['source_sha256'], 'generated_answers': False}, 'cells': cells}
            with output.open('x', encoding='utf-8') as stream:
                json.dump(notebook, stream, ensure_ascii=False, indent=2)
            validation = {'format': 'verified', 'source_pages': pages, 'content': 'extracted_text_requires_review', 'execution': 'not_run', 'answers': 'not_provided'}
        elif kind in {'docx', 'pdf'}:
            content = request['content']
            title = request.get('title', '')
            if isinstance(content, str):
                paragraphs = content.splitlines()
            elif isinstance(content, dict):
                title = content.get('title', title)
                paragraphs = content.get('paragraphs', [])
            else:
                raise ResourceError('INVALID_ARTIFACT', '文档正文需要文字或段落列表。')
            if not isinstance(paragraphs, list) or len(paragraphs) > 3000 or any(not isinstance(p, str) for p in paragraphs):
                raise ResourceError('INVALID_ARTIFACT', '文档段落格式无效或数量过多。')
            if sum(map(len, paragraphs)) > 128000:
                raise ResourceError('RESOURCE_LIMIT', '文档正文超过生成预算。')
            if kind == 'docx':
                from docx import Document
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn
                from docx.shared import Pt, Cm
                document = Document()
                section = document.sections[0]
                section.top_margin = section.bottom_margin = Cm(2)
                style = document.styles['Normal']
                style.font.name, style.font.size = 'Microsoft YaHei', Pt(11)
                style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
                if title:
                    document.add_heading(title, 0)
                for paragraph in paragraphs:
                    document.add_paragraph(paragraph)
                with output.open('xb') as stream:
                    document.save(stream)
                checked = Document(output)
                validation = {'format': 'verified', 'paragraphs': len(checked.paragraphs), 'content': 'not_checked', 'layout': 'not_rendered', 'execution': 'not_run'}
            else:
                from reportlab.pdfbase import pdfmetrics
                from reportlab.pdfbase.ttfonts import TTFont
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
                from reportlab.lib.pagesizes import A4
                font_path = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts' / 'simsun.ttc'
                font = 'Helvetica'
                if font_path.exists():
                    pdfmetrics.registerFont(TTFont('PMChinese', str(font_path), subfontIndex=0))
                    font = 'PMChinese'
                elif any(ord(ch) > 255 for ch in title + ''.join(paragraphs)):
                    raise ResourceError('FONT_REQUIRED', '需要可嵌入的中文字体才能生成此 PDF。')
                styles = getSampleStyleSheet()
                for style in styles.byName.values():
                    style.fontName = font
                    style.wordWrap = 'CJK'
                story = []
                if title:
                    story.extend([Paragraph(escape(title), styles['Title']), Spacer(1, 14)])
                for paragraph in paragraphs:
                    story.extend([Paragraph(escape(paragraph) or '&#160;', styles['BodyText']), Spacer(1, 6)])
                with output.open('xb') as stream:
                    SimpleDocTemplate(stream, pagesize=A4, topMargin=48, bottomMargin=48).build(story)
                from pypdf import PdfReader
                checked = PdfReader(output)
                validation = {'format': 'verified', 'pages': len(checked.pages), 'content': 'not_checked', 'layout': 'not_rendered', 'execution': 'not_run'}
        else:
            raise ResourceError('UNSUPPORTED_ARTIFACT', '此格式没有注册文档执行器。')
        result = {'ok': True, 'validation': validation}
    except Exception as error:
        result = {'ok': False, 'code': getattr(error, 'code', 'DOCUMENT_ERROR'), 'message': getattr(error, 'message', str(error))[:500]}
    (workspace / '.executor-result.json').write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
    return 0 if result.get('ok') else 1


if __name__ == '__main__':
    sys.exit(worker(sys.argv[1]))
